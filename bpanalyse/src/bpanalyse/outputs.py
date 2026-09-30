"""Owns every file written to the output directory: workbooks, graph.json, graph.graphml and run.log."""

from __future__ import annotations

import io
import json
import re
import sqlite3
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font

from bpanalyse.graph import EDGE_TYPES
from bpanalyse.metrics import METRIC_NAMES
from bpanalyse.model import Collision, Estate, key
from bpanalyse.query import DEP_COLUMNS, named_edges

FIXED = datetime(2000, 1, 1)
INVENTORY = {
    "release": ("id, name, created, exported_by, source_file, package_id", "name, id"),
    "process": ("id, name, type, release_id, page_count, stage_count, version, language, global_code_hash, global_code_line_count", "name, id"),
    "page": ("id, process_id, name, type, is_main, is_published, stage_count", "name, process_id, id"),
    "stage": ("id, process_id, page_id, name, type, x, y", "name, process_id, id"),
    "code_stage": ("c.stage_id, c.process_id, c.language, c.line_count, c.input_count, c.output_count, c.code_hash", None),
    "data_item": ("c.stage_id, c.process_id, c.datatype, c.exposure, c.has_initial_value, c.is_masked, c.is_encrypted", None),
    "calc_stage": ("c.stage_id, c.process_id, c.expression_length, c.literal_count", None),
    "app_element": ("id, object_id, name, element_type, parent_id, attribute_count", "name, id"),
    "work_queue": ("name, key_field, from_release", "name"),
    "credential_ref": ("name, from_release, member_count", "name"),
    "group": ("id, name, kind, is_default, member_count", "name, id"),
    "group_member": ("group_id, member_id", "group_id, member_id"),
    "web_api_service": ("id, name, enabled, action_count", "name, id"),
    "environment_variable": ("name, datatype, from_release", "name"),
    "external_object": ("name", "name"),
}
FINDING_COLUMNS = ("finding_type", "entity_name", "entity_type", "page_name", "stage_name", "detail")


def _pin_zip(data: bytes) -> bytes:
    """openpyxl stamps `modified` and every zip entry with the current time; rewrite them to the fixed value."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                body = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + FIXED.strftime("%Y-%m-%dT%H:%M:%SZ").encode(), body)
            entry = zipfile.ZipInfo(info.filename, (2000, 1, 1, 0, 0, 0))
            entry.compress_type, entry.external_attr = zipfile.ZIP_DEFLATED, 0o600 << 16
            dst.writestr(entry, body)
    return out.getvalue()


def write_xlsx(path: Path, sheets: list[tuple[str, tuple[str, ...], list[tuple]]]) -> int:
    """Write the workbook byte-deterministically; return the number of illegal control characters stripped from strings."""
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.created = wb.properties.modified = FIXED
    stripped = 0
    for title, columns, rows in sheets:
        ws = wb.create_sheet(title)
        for r, row in enumerate([columns, *rows], start=1):
            clean = []
            for v in row:
                if isinstance(v, str):
                    v, n = ILLEGAL_CHARACTERS_RE.subn("", v)
                    stripped += n
                clean.append(v)
            ws.append(clean)
            for cell in ws[r]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # a name such as "=SUM(A1)" must stay text, not become a formula
                if r == 1:
                    cell.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    path.write_bytes(_pin_zip(buf.getvalue()))
    return stripped


def _select(conn: sqlite3.Connection, sql: str) -> list[tuple]:
    return [tuple(r) for r in conn.execute(sql)]


def inventory_sheets(conn: sqlite3.Connection) -> list[tuple[str, tuple[str, ...], list[tuple]]]:
    sheets = []
    for table, (cols, order) in INVENTORY.items():
        if order is None:
            sql = f"SELECT {cols} FROM \"{table}\" c JOIN stage s ON s.id = c.stage_id AND s.process_id = c.process_id ORDER BY s.name, c.process_id, c.stage_id"
        else:
            sql = f'SELECT {cols} FROM "{table}" ORDER BY {order}'
        names = tuple(c.strip().split(".")[-1] for c in cols.split(","))
        sheets.append((table, names, _select(conn, sql)))
    return sheets


def dependency_sheets(conn: sqlite3.Connection) -> list[tuple[str, tuple[str, ...], list[tuple]]]:
    edges = named_edges(conn)
    return [(t, DEP_COLUMNS, [e.row for e in edges if e.edge_type == t]) for t in EDGE_TYPES]


def metric_rows(conn: sqlite3.Connection) -> list[tuple]:
    cols = ", ".join(METRIC_NAMES)
    sql = f"SELECT name, type, {cols} FROM metrics ORDER BY fan_in DESC, name, id"
    return [tuple(int(v) if isinstance(v, float) and v.is_integer() else v for v in r) for r in conn.execute(sql)]


def graph_data(conn: sqlite3.Connection) -> dict:
    metrics = {i: dict(zip(METRIC_NAMES, [int(v) if float(v).is_integer() else v for v in r]))
               for i, *r in conn.execute(f"SELECT id, {', '.join(METRIC_NAMES)} FROM metrics")}
    nodes: dict[tuple[str, str], dict] = {}
    for r in conn.execute("SELECT id, name, type, release_id, page_count, stage_count, version, language, global_code_hash, "
                          "global_code_line_count FROM process"):
        nodes[(r[2], r[0])] = {"id": r[0], "type": r[2], "name": r[1], "release_id": r[3], "page_count": r[4],
                               "stage_count": r[5], "version": r[6], "language": r[7], "global_code_hash": r[8],
                               "global_code_line_count": r[9], "metrics": metrics.get(r[0], {}), "in_estate": True}
    for r in conn.execute("SELECT id, process_id, name, type, is_main, is_published, stage_count FROM page"):
        nodes[("page", key(r[1], r[0]))] = {"id": key(r[1], r[0]), "type": "page", "name": r[2], "process_id": r[1], "page_type": r[3],
                                 "is_main": r[4], "is_published": r[5], "stage_count": r[6], "in_estate": True}
    for r in conn.execute("SELECT id, object_id, name, element_type, parent_id, attribute_count FROM app_element"):
        nodes[("app_element", key(r[1], r[0]))] = {"id": key(r[1], r[0]), "type": "app_element", "name": r[2], "object_id": r[1],
                                        "element_type": r[3], "parent_id": r[4], "attribute_count": r[5], "in_estate": True}
    for table, kind, extra in (("work_queue", "work_queue", "key_field"), ("credential_ref", "credential_ref", "from_release"),
                               ("environment_variable", "environment_variable", "datatype"), ("external_object", "external_object", "name")):
        for r in conn.execute(f"SELECT name, {extra} FROM {table}"):
            nodes[(kind, r[0])] = {"id": r[0], "type": kind, "name": r[0], extra: r[1], "in_estate": kind != "external_object"}
    cols = ("from_id", "from_process_id", "from_type", "to_id", "to_type", "edge_type", "evidence_stage_id", "evidence_process_id",
            "unresolved", "release_id", "detail", "is_internal")
    edges = [dict(zip(cols, r)) for r in conn.execute(
        f"SELECT {', '.join(cols)} FROM edge ORDER BY edge_type, from_id, to_id, evidence_process_id, evidence_stage_id")]
    for e in edges:
        if e["to_id"] and (e["to_type"], e["to_id"]) not in nodes:
            nodes[(e["to_type"], e["to_id"])] = {"id": e["to_id"], "type": e["to_type"], "name": e["to_id"], "in_estate": False}
    return {"nodes": [nodes[k] for k in sorted(nodes)], "edges": edges}


def write_graph_json(path: Path, data: dict) -> None:
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(json.dumps(data, sort_keys=True, indent=2) + "\n")


def write_graphml(path: Path, data: dict) -> None:
    ns = "http://graphml.graphdrawing.org/xmlns"
    ET.register_namespace("", ns)
    root = ET.Element(f"{{{ns}}}graphml")
    for kid, target, name in (("d0", "node", "type"), ("d1", "node", "name"), ("d2", "edge", "edge_type"),
                              ("d3", "edge", "evidence_stage_id"), ("d4", "edge", "unresolved")):
        ET.SubElement(root, f"{{{ns}}}key", {"id": kid, "for": target, "attr.name": name, "attr.type": "string"})
    g = ET.SubElement(root, f"{{{ns}}}graph", {"id": "estate", "edgedefault": "directed"})
    for n in data["nodes"]:
        el = ET.SubElement(g, f"{{{ns}}}node", {"id": n["type"] + ":" + n["id"]})
        for k, v in (("d0", n["type"]), ("d1", n["name"])):
            ET.SubElement(el, f"{{{ns}}}data", {"key": k}).text = v
    for i, e in enumerate(x for x in data["edges"] if x["to_id"]):
        el = ET.SubElement(g, f"{{{ns}}}edge", {"id": "e" + str(i), "source": e["from_type"] + ":" + e["from_id"],
                                              "target": e["to_type"] + ":" + e["to_id"]})
        for k, v in (("d2", e["edge_type"]), ("d3", e["evidence_stage_id"]), ("d4", str(e["unresolved"]))):
            ET.SubElement(el, f"{{{ns}}}data", {"key": k}).text = v
    ET.indent(root)
    ET.ElementTree(root).write(path, encoding="utf-8", xml_declaration=True)


def build_log(files: list[str], per_file: dict[str, Counter], edge_counts: Counter, unknowns: Counter,
              collisions: list[Collision], estate: Estate, failures: list[tuple[str, int, str, str]], notes: list[str]) -> list[str]:
    lines = ["file_count=" + str(len(files) + len(failures))]
    for name in sorted(per_file):
        lines.append("file=" + name + " " + " ".join(f"{k}={v}" for k, v in sorted(per_file[name].items())))
    lines += [f"edges type={t} count={edge_counts[t]}" for t in EDGE_TYPES]
    lines += [f"unrecognised {path} {n}" for path, n in sorted(unknowns.items())]
    lines += sorted(notes)
    for c in sorted(collisions, key=lambda c: (c.name, c.id, c.dropped_file)):
        lines.append(f"collision id={c.id} name=\"{c.name}\" kept={c.kept_file} kept_created={c.kept_created} "
                     f"dropped={c.dropped_file} dropped_created={c.dropped_created}")
    masked = [(pd.process.name, pd.masked_total) for pd in estate.processes if pd.masked_total]
    masked += [(n, c) for n, (_, c) in estate.env_masks.items()]
    lines += [f"masked owner=\"{n}\" count={c}" for n, c in sorted(masked)]
    lines.append("masked_total=" + str(sum(c for _, c in masked)))
    for f, ln, event, fields in sorted(failures):
        if event != "parse_error":
            lines.append(f"{event} file={f} {fields}")
        lines.append(f"parse_error file={f} line={ln}" + (f" {fields}" if event == "parse_error" and fields else ""))
    return lines


def write_all(conn: sqlite3.Connection, out: Path, log: list[str]) -> None:
    stripped = write_xlsx(out / "inventory.xlsx", inventory_sheets(conn))
    stripped += write_xlsx(out / "dependencies.xlsx", dependency_sheets(conn))
    stripped += write_xlsx(out / "release_map.xlsx", [("release_map", ("process_name", "object_name", "action_name", "call_count"),
        _select(conn, "SELECT * FROM release_map ORDER BY process_name, object_name, action_name"))])
    stripped += write_xlsx(out / "metrics.xlsx", [("metrics", ("name", "type", *METRIC_NAMES), metric_rows(conn))])
    stripped += write_xlsx(out / "findings.xlsx", [("findings", FINDING_COLUMNS, _select(conn,
        "SELECT * FROM finding ORDER BY finding_type, entity_name, page_name, stage_name, detail"))])
    data = graph_data(conn)
    write_graph_json(out / "graph.json", data)
    write_graphml(out / "graph.graphml", data)
    log = [*log, f"illegal_chars_stripped={stripped}"]
    (out / "run.log").write_text("\n".join(log) + "\n", encoding="utf-8", newline="\n")

"""SPEC 2.6, 4.7, 8.2: odd-but-valid release layouts parse; unusable ones fail the file with a located log line."""

import sqlite3

import pytest
from conftest import MAIN, NS, process_xml, release_xml, run_files, stage_xml, wrapper_xml
from openpyxl import load_workbook

from bpanalyse.cli import main
from bpanalyse.metrics import _depths
from bpanalyse.outputs import write_xlsx

PID = "22222222-0000-0000-0000-000000000002"


def _names(out):
    return {r[0] for r in sqlite3.connect(out / "estate.sqlite").execute("SELECT name FROM process")}


def _log(out):
    return (out / "run.log").read_text().splitlines()


def test_bom_crlf_declaration_and_prefixed_inner(tmp_path):
    inner = '<?xml version="1.0" encoding="utf-16"?>\r\n﻿' + process_xml(PID, "Escaped", prefix="x:").replace("<x:process", f'<x:process xmlns:x="{NS}"', 1)
    text = release_xml(wrapper_xml(PID, "Escaped", inner, escape=True), eol="\r\n")
    rc, out = run_files(tmp_path, {"a.bprelease": b"\xef\xbb\xbf" + text.encode("utf-8")})
    assert rc == 0 and _names(out) == {"Escaped"}


def test_default_namespace_file(tmp_path):
    text = release_xml(wrapper_xml(PID, "Plain", process_xml(PID, "Plain"))).replace("bpr:", "").replace(f' xmlns:bpr="{NS}"', f' xmlns="{NS}"')
    assert 'xmlns="' in text and "<bpr:" not in text
    rc, out = run_files(tmp_path, {"a.bprelease": text})
    assert rc == 0 and _names(out) == {"Plain"}


def test_unparsed_wrapper_fails_file(tmp_path):
    rc, out = run_files(tmp_path, {"a.bprelease": release_xml(wrapper_xml(PID, "Hollow", "no process here")),
                                   "b.bprelease": release_xml(wrapper_xml("33", "Fine", process_xml("33", "Fine")), pkg="p2")})
    log = _log(out)
    assert rc == 2 and _names(out) == {"Fine"}
    assert any(line.startswith("unparsed_wrapper file=a.bprelease") and f"id={PID}" in line and '"Hollow"' in line for line in log)
    assert any(line.startswith("parse_error file=a.bprelease line=") for line in log)


def test_inner_parse_error_reports_wrapper_line_plus_offset(tmp_path):
    broken = process_xml(PID, "Broken").replace("</process>", "</wrong>")
    text = release_xml(wrapper_xml(PID, "Broken", "\n\n" + broken, escape=True), manifest=f'<bpr:process id="{PID}" name="Broken"/>')
    rc, out = run_files(tmp_path, {"a.bprelease": text})
    wrapper_line = next(i for i, line in enumerate(text.splitlines(), 1) if line.startswith("<bpr:process"))
    (entry,) = [line for line in _log(out) if line.startswith("parse_error")]
    assert rc == 2 and int(entry.rsplit("line=", 1)[1]) == wrapper_line + 2


def test_missing_process_id_uses_manifest_then_fails(tmp_path):
    manifest = f'<bpr:process id="{PID}" name="Anon"/>'
    good = release_xml(wrapper_xml(None, "Anon", process_xml(None, "Anon")), manifest=manifest)
    rc, out = run_files(tmp_path, {"a.bprelease": good})
    assert rc == 0 and sqlite3.connect(out / "estate.sqlite").execute("SELECT id FROM process").fetchall() == [(PID,)]
    bad = release_xml(wrapper_xml(None, "Nameless", process_xml(None, "Nameless")))
    rc, out = run_files(tmp_path / "x", {"a.bprelease": bad}) if (tmp_path / "x").mkdir() is None else (0, tmp_path)
    assert rc == 2 and any(line.startswith("parse_error file=a.bprelease line=") and 'process="Nameless"' in line for line in _log(out))


def test_bad_created_logged_without_value(tmp_path):
    a = release_xml(wrapper_xml(PID, "P", process_xml(PID, "P")), created="not-a-date-SECRET")
    b = release_xml(wrapper_xml(PID, "P", process_xml(PID, "P")), created="2026-02-01T00:00:00Z", pkg="p2")
    rc, out = run_files(tmp_path, {"a.bprelease": a, "b.bprelease": b})
    log = "\n".join(_log(out))
    assert rc == 0 and "bad_created file=a.bprelease value_length=17" in log
    assert "dropped_created=unparsed" in log and "SECRET" not in log


def test_ambiguity_and_missing_reference_notes(tmp_path):
    pages = MAIN + (("pg-a", "Do It", "Normal", "True"), ("pg-b", "Do It", "Normal", "True"))
    obj = lambda i: wrapper_xml(i, "Twin", process_xml(i, "Twin", pages=pages, obj=True), obj=True)  # noqa: E731
    caller = process_xml("cc", "Caller", [stage_xml("s1", "Act", "Action", inner="<inputs/>"),
                                          stage_xml("s2", "Call", "Process", inner="<target/>")])
    rc, out = run_files(tmp_path, {"a.bprelease": release_xml(obj("t1"), obj("t2"), wrapper_xml("cc", "Caller", caller))})
    log = _log(out)
    assert rc == 0
    assert 'ambiguous_object name="Twin" ids=t1,t2' in log
    assert 'ambiguous_action object="Twin" action="Do It" count=4' in log
    assert 'stage_missing_resource process="Caller" stage="Act"' in log
    assert 'stage_missing_resource process="Caller" stage="Call"' in log
    assert sqlite3.connect(out / "estate.sqlite").execute("SELECT count(*) FROM edge WHERE edge_type = 'calls_process'").fetchone() == (0,)


def test_empty_input_dir_and_unknown_query_name(tmp_path, capsys):
    (tmp_path / "in").mkdir()
    assert main(["run", str(tmp_path / "in"), str(tmp_path / "o")]) == 1
    assert "no_input_files dir=" in capsys.readouterr().err
    rc, out = run_files(tmp_path, {"UPPER.BPRELEASE": release_xml(wrapper_xml(PID, "P", process_xml(PID, "P")))})
    assert rc == 0 and _names(out) == {"P"}
    assert main(["query", "impact", "Ghost", "--db", str(out / "estate.sqlite")]) == 1
    assert "not_found name=Ghost" in capsys.readouterr().err
    assert main(["query", "impact", "P", "--db", str(out / "estate.sqlite")]) == 0


def test_xlsx_cells_are_text_and_illegal_characters_stripped(tmp_path):
    path = tmp_path / "t.xlsx"
    n = write_xlsx(path, [("s", ("name",), [("=SUM(A1)",), ("a\x01b\x02",), (7,)])])
    ws = load_workbook(path)["s"]
    assert n == 2 and [c.value for c in ws["A"]] == ["name", "=SUM(A1)", "ab", 7]
    assert ws["A2"].data_type == "s"


def test_depth_counts_a_cycle_once_and_is_not_exponential():
    assert _depths({"a": {"b"}, "b": {"c"}, "c": {"b", "d"}}) == {"a": 2, "b": 1, "c": 1, "d": 0}
    layers = {f"n{i}_{j}": {f"n{i + 1}_{k}" for k in range(6)} for i in range(40) for j in range(6)}
    assert _depths(layers)["n0_0"] == 40
    assert _depths({f"x{i}": {f"x{i + 1}"} for i in range(5000)})["x0"] == 5000
